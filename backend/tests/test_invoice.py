import io

from openpyxl import load_workbook

from app.models import UserRole
from tests.test_orders import CHECKOUT


def rows(content: bytes) -> list[tuple]:
    ws = load_workbook(io.BytesIO(content)).active
    return list(ws.iter_rows(values_only=True))


def test_invoice_lists_what_is_shipped(client, catalog, auth):
    admin = auth(UserRole.admin)
    order = client.post(
        "/api/orders",
        json={
            **CHECKOUT,
            "items": [
                {"variant_id": catalog["coco50"].id, "quantity": 3},
                {"variant_id": catalog["coco100"].id, "quantity": 1},
            ],
        },
        headers=auth(UserRole.retail),
    ).json()
    # The manager removes 100 ml before shipping: it must not be on the note.
    line100 = next(i for i in order["items"] if i["volume_ml"] == 100)
    r = client.patch(
        f"/api/admin/orders/{order['id']}/items",
        json={"items": [{"order_item_id": line100["id"], "quantity": 0}], "reason": "нет"},
        headers=admin,
    )
    assert r.status_code == 200, r.text

    r = client.get(f"/api/admin/orders/{order['id']}/invoice.xlsx", headers=admin)
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml"
    )
    assert f"nakladnaya_{order['id']}.xlsx" in r.headers["content-disposition"]

    table = rows(r.content)
    assert table[0][0].startswith(f"Накладная № {order['id']} от ")
    flat = [v for row in table for v in row if v is not None]
    assert "Анна" in flat and "Алматы, ул. Абая, 1" in flat
    header = next(row for row in table if row[0] == "№")
    assert header[3:7] == ("Тип", "Пол", "Объём, мл", "Тестер")
    items = [row for row in table if row[0] == 1]
    assert items == [
        (1, "Chanel", "Coco Mademoiselle", "EDP", "Женский", 50, "Нет", None, 3, 100, 300)
    ]
    total = next(row for row in table if "Итого" in row)
    assert (total[8], total[10]) == (3, 300)

    # The admin order has the same product details for the printable note.
    admin_order = client.get(f"/api/admin/orders/{order['id']}", headers=admin).json()
    line = next(i for i in admin_order["items"] if i["volume_ml"] == 50)
    assert (line["product_type"], line["gender"], line["is_tester"]) == ("EDP", "female", False)


def test_invoice_is_admin_only(client, catalog, auth):
    order = client.post(
        "/api/orders",
        json={**CHECKOUT, "items": [{"variant_id": catalog["coco50"].id, "quantity": 1}]},
        headers=auth(UserRole.retail),
    ).json()
    url = f"/api/admin/orders/{order['id']}/invoice.xlsx"
    assert client.get(url).status_code == 401
    assert client.get(url, headers=auth(UserRole.retail)).status_code == 403
    missing = client.get("/api/admin/orders/999999/invoice.xlsx", headers=auth(UserRole.admin))
    assert missing.status_code == 404
