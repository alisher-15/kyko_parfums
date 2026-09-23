"""Import brands / products / variants from an Excel (.xlsx) or CSV price list.

The importer is tolerant to the layout of the source file:

* the header row is searched among the first rows of the sheet;
* column names are matched against Russian/English aliases (see ``COLUMN_ALIASES``);
* one row = one product, optionally with one variant (volume + prices + stock).
  Several rows with the same brand/name/type add several volumes to the same product;
* re-importing the same file is idempotent: products are matched by (brand, name, type) and
  only non-empty cells overwrite existing values.
"""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload, selectinload

from app.models import Brand, Gender, Product, ProductVariant
from app.schemas.admin import ImportReport, ImportRowError, check_price_order

# fmt: off
COLUMN_ALIASES: dict[str, list[str]] = {
    "brand": ["brand", "бренд", "марка", "производитель", "торговая марка"],
    "name": ["name", "название", "наименование", "аромат", "товар", "product", "название аромата"],
    "type": ["type", "тип", "концентрация", "вид", "тип аромата"],
    "category": [
        "category", "категория", "олфактивная группа", "группа", "семейство",
        "группа ароматов", "olfactory group", "family",
    ],
    "gender": ["gender", "пол", "для кого", "sex"],
    "longevity": ["longevity", "стойкость"],
    "top_notes": [
        "top_notes", "top notes", "верхние ноты", "ноты верха", "верх", "начальные ноты",
        "ноты верхние",
    ],
    "mid_notes": [
        "mid_notes", "middle notes", "heart notes", "ноты сердца", "средние ноты", "сердце",
        "ноты средние", "ноты сердце",
    ],
    "base_notes": [
        "base_notes", "base notes", "базовые ноты", "ноты базы", "база", "конечные ноты",
        "ноты базовые", "ноты база", "шлейф",
    ],
    "description": ["description", "описание"],
    "image_url": ["image_url", "image", "фото", "photo", "photo_url", "изображение", "картинка"],
    "volume_ml": ["volume_ml", "volume", "объем", "объем мл", "мл", "объем ml"],
    "retail_price": [
        "retail_price", "розница", "розничная цена", "цена розница", "цена", "price",
        "цена розничная",
    ],
    "wholesale_price": ["wholesale_price", "опт", "оптовая цена", "цена опт", "цена оптовая"],
    "bulk_price": [
        "bulk_price", "крупный опт", "цена крупный опт", "цена крупного опта", "крупнооптовая цена",
    ],
    "stock": ["stock", "остаток", "остатки", "количество", "кол во", "наличие", "склад"],
    "sku": ["sku", "артикул", "код", "штрихкод", "barcode"],
    "is_active": ["is_active", "активен", "активный", "опубликован"],
}
# fmt: on

TEMPLATE_COLUMNS = [
    ("brand", "Бренд"),
    ("name", "Название"),
    ("type", "Тип"),
    ("category", "Олфактивная группа"),
    ("gender", "Пол"),
    ("longevity", "Стойкость"),
    ("top_notes", "Верхние ноты"),
    ("mid_notes", "Ноты сердца"),
    ("base_notes", "Базовые ноты"),
    ("description", "Описание"),
    ("image_url", "Фото"),
    ("volume_ml", "Объём, мл"),
    ("retail_price", "Розничная цена"),
    ("wholesale_price", "Оптовая цена"),
    ("bulk_price", "Цена крупный опт"),
    ("stock", "Остаток"),
    ("sku", "Артикул"),
]

PRODUCT_TEXT_FIELDS = (
    "category",
    "longevity",
    "top_notes",
    "mid_notes",
    "base_notes",
    "description",
    "image_url",
)
HEADER_SCAN_ROWS = 15


def _norm_header(value: Any) -> str:
    s = str(value or "").strip().lower().replace("ё", "е").replace("_", " ")
    s = re.sub(r"[^\w\s]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


_ALIAS_LOOKUP = {
    _norm_header(alias): key for key, aliases in COLUMN_ALIASES.items() for alias in aliases
}


def _text(value: Any, max_len: int | None = None) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = re.sub(r"\s+", " ", str(value)).strip()
    if not s or s in {"-", "—", "nan", "None"}:
        return None
    return s[:max_len] if max_len else s


def normalize_gender(value: Any) -> Gender | None:
    s = (_text(value) or "").lower().replace("ё", "е")
    if not s:
        return None
    if "уни" in s or "unisex" in s or s in {"u", "у"}:
        return Gender.unisex
    if s.startswith("жен") or "women" in s or "woman" in s or "female" in s or s in {"ж", "f", "w"}:
        return Gender.female
    if s.startswith("муж") or "men" in s or "man" in s or "male" in s or s in {"м", "m"}:
        return Gender.male
    raise ValueError(f"не удалось распознать пол «{value}»")


def normalize_type(value: Any) -> str | None:
    s = _text(value)
    if s is None:
        return None
    low = s.lower().replace("ё", "е").replace(".", " ")
    low = re.sub(r"\s+", " ", low).strip()
    rules = [
        (("edp", "eau de parfum", "парфюмерная вода", "парф вода"), "EDP"),
        (("edt", "eau de toilette", "туалетная вода", "туал вода"), "EDT"),
        (("edc", "eau de cologne", "одеколон", "cologne"), "EDC"),
        (("extrait", "экстракт", "extract"), "Extrait"),
        (("parfum", "perfume", "духи", "парфюм"), "Parfum"),
    ]
    for needles, canonical in rules:
        if any(n == low or n in low for n in needles):
            return canonical
    return s[:32]


def parse_decimal(value: Any) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    s = re.sub(r"[^\d.,-]", "", str(value))
    if not s or s in {"-", ".", ","}:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if re.fullmatch(r"-?\d+,\d{1,2}", s) else s.replace(",", "")
    try:
        d = Decimal(s).quantize(Decimal("0.01"))
    except InvalidOperation as e:
        raise ValueError(f"не число: «{value}»") from e
    if d < 0:
        raise ValueError(f"отрицательное значение: «{value}»")
    return d


def parse_int(value: Any) -> int | None:
    d = parse_decimal(value)
    if d is None:
        return None
    if d != d.to_integral_value():
        raise ValueError(f"ожидалось целое число: «{value}»")
    return int(d)


def parse_bool(value: Any) -> bool | None:
    s = (_text(value) or "").lower()
    if not s:
        return None
    if s in {"1", "да", "yes", "true", "y", "д", "+"}:
        return True
    if s in {"0", "нет", "no", "false", "n", "н", "-"}:
        return False
    raise ValueError(f"не удалось распознать да/нет: «{value}»")


# ---------- Reading ----------


def read_rows(
    content: bytes, filename: str, sheet: str | None = None
) -> tuple[list[tuple[int, list[Any]]], int]:
    """Return ([(excel_row_number, cells)], header_row_index_in_list)."""
    if filename.lower().endswith(".csv"):
        text = _decode(content)
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        rows = list(csv.reader(io.StringIO(text), dialect))
    else:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    numbered = [(i + 1, r) for i, r in enumerate(rows)]

    for idx, (_, cells) in enumerate(numbered[:HEADER_SCAN_ROWS]):
        keys = {_ALIAS_LOOKUP.get(_norm_header(c)) for c in cells}
        if {"brand", "name"} <= keys:
            return numbered, idx
    raise ValueError("Не найдена строка заголовков: нужны как минимум колонки «Бренд» и «Название»")


def _decode(content: bytes) -> str:
    for enc in ("utf-8-sig", "cp1251"):
        try:
            return content.decode(enc)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


# ---------- Import ----------


@dataclass
class _Stats:
    rows_total: int = 0
    rows_skipped: int = 0
    brands_created: int = 0
    products_created: int = 0
    products_updated: int = 0
    variants_created: int = 0
    variants_updated: int = 0
    errors: list[ImportRowError] = field(default_factory=list)


def import_catalog(
    db: Session,
    content: bytes,
    filename: str,
    sheet: str | None = None,
    dry_run: bool = False,
) -> ImportReport:
    rows, header_idx = read_rows(content, filename, sheet)
    header = rows[header_idx][1]
    columns: dict[str, int] = {}
    unmapped: list[str] = []
    for i, cell in enumerate(header):
        key = _ALIAS_LOOKUP.get(_norm_header(cell))
        if key and key not in columns:
            columns[key] = i
        elif _text(cell):
            unmapped.append(str(cell).strip())

    brands = {b.name.lower(): b for b in db.scalars(select(Brand))}
    products: dict[tuple[str, str, str], Product] = {}
    for p in db.scalars(
        select(Product).options(joinedload(Product.brand), selectinload(Product.variants))
    ).unique():
        products[(p.brand.name.lower(), p.name.lower(), (p.type or "").lower())] = p
    sku_owner: dict[str, ProductVariant] = {
        v.sku: v for v in db.scalars(select(ProductVariant).where(ProductVariant.sku.is_not(None)))
    }
    touched_products: set[int] = set()  # id() of products updated in this run

    st = _Stats()

    def cell(cells: list[Any], key: str) -> Any:
        i = columns.get(key)
        return cells[i] if i is not None and i < len(cells) else None

    for row_no, cells in rows[header_idx + 1 :]:
        if not any(_text(c) for c in cells):
            continue
        st.rows_total += 1
        try:
            brand_name = _text(cell(cells, "brand"), 255)
            name = _text(cell(cells, "name"), 255)
            if not brand_name or not name:
                raise ValueError("не заполнены бренд или название")

            ptype = normalize_type(cell(cells, "type"))
            values = {f: _text(cell(cells, f)) for f in PRODUCT_TEXT_FIELDS}
            if values["category"]:
                values["category"] = values["category"][:128]
            if values["longevity"]:
                values["longevity"] = values["longevity"][:64]
            gender = normalize_gender(cell(cells, "gender"))
            is_active = parse_bool(cell(cells, "is_active"))

            volume = parse_int(cell(cells, "volume_ml"))
            retail = parse_decimal(cell(cells, "retail_price"))
            wholesale = parse_decimal(cell(cells, "wholesale_price"))
            bulk = parse_decimal(cell(cells, "bulk_price"))
            stock = parse_int(cell(cells, "stock"))
            sku = _text(cell(cells, "sku"), 64)
            if volume is not None and volume <= 0:
                raise ValueError("объём должен быть больше нуля")
        except ValueError as e:
            st.errors.append(ImportRowError(row=row_no, error=str(e)))
            st.rows_skipped += 1
            continue

        # Brand
        brand = brands.get(brand_name.lower())
        if brand is None:
            brand = Brand(name=brand_name)
            db.add(brand)
            brands[brand_name.lower()] = brand
            st.brands_created += 1

        # Product
        key = (brand_name.lower(), name.lower(), (ptype or "").lower())
        product = products.get(key)
        if product is None:
            product = Product(brand=brand, name=name, type=ptype, gender=gender, **values)
            if is_active is not None:
                product.is_active = is_active
            db.add(product)
            products[key] = product
            touched_products.add(id(product))
            st.products_created += 1
        else:
            changed = False
            for f, v in {**values, "gender": gender, "is_active": is_active}.items():
                if v is not None and getattr(product, f) != v:
                    setattr(product, f, v)
                    changed = True
            if changed and id(product) not in touched_products:
                touched_products.add(id(product))
                st.products_updated += 1

        # Variant
        if volume is None:
            continue
        variant = next((v for v in product.variants if v.volume_ml == volume), None)
        if sku and sku in sku_owner and sku_owner[sku] is not variant:
            st.errors.append(
                ImportRowError(row=row_no, error=f"артикул {sku} уже занят — артикул не сохранён")
            )
            sku = None
        if variant is None:
            try:
                check_price_order(retail, wholesale, bulk)
            except ValueError as e:
                st.errors.append(ImportRowError(row=row_no, error=f"объём {volume} мл: {e}"))
                continue
            if retail is None:
                st.errors.append(
                    ImportRowError(
                        row=row_no,
                        error=f"объём {volume} мл пропущен: не указана розничная цена",
                    )
                )
                continue
            variant = ProductVariant(
                volume_ml=volume,
                retail_price=retail,
                wholesale_price=wholesale,
                bulk_price=bulk,
                stock=stock or 0,
                sku=sku,
            )
            product.variants.append(variant)
            st.variants_created += 1
        else:
            updates = {
                "retail_price": retail,
                "wholesale_price": wholesale,
                "bulk_price": bulk,
                "stock": stock,
                "sku": sku,
            }
            new = {k: v for k, v in updates.items() if v is not None}
            merged = {k: new.get(k, getattr(variant, k)) for k in updates}
            try:
                check_price_order(
                    merged["retail_price"], merged["wholesale_price"], merged["bulk_price"]
                )
            except ValueError as e:
                st.errors.append(ImportRowError(row=row_no, error=f"объём {volume} мл: {e}"))
                continue
            if any(getattr(variant, k) != v for k, v in new.items()):
                for k, v in new.items():
                    setattr(variant, k, v)
                st.variants_updated += 1
        if sku:
            sku_owner[sku] = variant

    if dry_run:
        db.rollback()
    else:
        db.commit()

    return ImportReport(
        dry_run=dry_run,
        rows_total=st.rows_total,
        rows_skipped=st.rows_skipped,
        brands_created=st.brands_created,
        products_created=st.products_created,
        products_updated=st.products_updated,
        variants_created=st.variants_created,
        variants_updated=st.variants_updated,
        errors=st.errors,
        unmapped_columns=unmapped,
    )


def build_template() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Каталог"
    ws.append([title for _, title in TEMPLATE_COLUMNS])
    # fmt: off
    ws.append(
        [
            "Chanel", "Coco Mademoiselle", "EDP", "Шипровые", "Женский", "Стойкий",
            "Апельсин, бергамот", "Роза, жасмин", "Пачули, ветивер", "Описание аромата", "",
            50, 65000, 55000, 50000, 10, "CH-CM-50",
        ]
    )
    ws.append(
        [
            "Chanel", "Coco Mademoiselle", "EDP", "", "", "", "", "", "", "", "",
            100, 95000, 82000, 76000, 5, "CH-CM-100",
        ]
    )
    # fmt: on
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
