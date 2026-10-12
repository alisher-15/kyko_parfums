import io
from decimal import Decimal

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import event, select

from app.db import engine
from app.models import (
    PriceFormat,
    ProductVariant,
    Supplier,
    SupplierLink,
    SupplierPrice,
    UserRole,
    VariantBarcode,
)
from app.services import markups
from app.services.price_lists import PriceListError, read_price_list
from app.services.suppliers import import_links, price_from_cost


def _xlsx(rows: list[list]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def kanat_file(items: list[tuple[str, float]]) -> bytes:
    """Канат's 1С layout: (line, price in tenge); the dollar price is tenge / 495."""
    return _xlsx(
        [
            [None, "Прайс-лист"],
            [None, "Цены указаны на 23.09.2026"],
            [None, "Ценовая группа/ Номенклатура/ Характеристика номенклатуры", "Долларовая"],
            [None, None, "Цена", "Ед."],
            [None, "1. Элитка", None, None, "ОПТ,ЦЕНА", "ЗАКАЗЫ", "СУММА"],
            *([None, text, round(kzt / 495, 2), "шт", kzt, None, 0] for text, kzt in items),
            [None, "2. Уход", None, None, 0, None, 0],
        ]
    )


def erik_file(items: list[tuple[str | None, str, float]]) -> bytes:
    """Эрик's layout: (barcode, line, price); the cash price next to it is not used."""
    return _xlsx(
        [
            ["опт ", None, "Товар обмену,возврату не подлежит.", "ПРИМ", "ОПТ ", 480, "ОПЛАТА"],
            [480, None, None, None, "СКИДКА ПО ПРЕДОПЛАТЕ НАЛИЧНЫМИ % ", 3, "К ОПЛАТЕ"],
            [None, None, "рабочие дни: вторник, среда"],
            [None, "ШТРИХ - КОД"],
            ["15,,09", "15,,09", "обработка заказа 2 рабочих дня", None, "15,,09", 0, 0],
            *([kzt, code, text, None, kzt, round(kzt * 0.97)] for code, text, kzt in items),
        ]
    )


@pytest.fixture
def suppliers(db):
    kanat = Supplier(name="Канат", is_main=True, price_format=PriceFormat.usd_1c)
    erik = Supplier(name="Эрик", price_format=PriceFormat.barcode_kzt)
    db.add_all([kanat, erik])
    db.commit()
    return {"kanat": kanat, "erik": erik}


def _priced(db, variant: ProductVariant, cost: str, retail: str, wholesale=None, bulk=None):
    """Prices set by hand: the volume remembers their markups."""
    variant.cost_price = Decimal(cost)
    variant.retail_price = Decimal(retail)
    variant.wholesale_price = Decimal(wholesale) if wholesale else None
    variant.bulk_price = Decimal(bulk) if bulk else None
    markups.remember(variant)
    db.commit()


def _upload(client, h, supplier, content: bytes, apply: bool = False, name: str = "price.xlsx"):
    files = {"file": (name, content, "application/octet-stream")}
    url = f"/api/admin/suppliers/{supplier.id}/price-list?dry_run={str(not apply).lower()}"
    r = client.post(url, files=files, headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def _link(client, h, supplier, line: str, variant_id: int | None):
    r = client.post(
        f"/api/admin/suppliers/{supplier.id}/links",
        json={"line": line, "variant_id": variant_id},
        headers=h,
    )
    assert r.status_code == 200, r.text


def _state(db, variant_id: int) -> tuple:
    db.expire_all()
    v = db.get(ProductVariant, variant_id)
    return (v.cost_price, v.retail_price, v.wholesale_price, v.bulk_price)


def test_price_from_cost_rounds_up_to_hundreds():
    assert price_from_cost(Decimal("11000"), Decimal("60")) == 17600
    assert price_from_cost(Decimal("10000"), Decimal("59.9088")) == 16000
    # 9999.45 × 1.599088 is 15 990.00 in whole tenge, not a hair above it.
    assert price_from_cost(Decimal("9999.45"), Decimal("59.9088")) == 16000


def test_reading_both_formats():
    lines = read_price_list(
        PriceFormat.usd_1c, kanat_file([("AFNAN 9 PM U EdP 100", 13860)]), "230926.xlsx"
    )
    assert [(x.text, x.price, x.price_usd) for x in lines] == [
        ("AFNAN 9 PM U EdP 100", Decimal("13860"), Decimal("28"))
    ]
    lines = read_price_list(
        PriceFormat.barcode_kzt,
        erik_file([(" 6290171002338", "AFNAN  9 PM   (U) edp   100ML", 13860), (None, "X", 500)]),
        "1509.xlsx",
    )
    # The date row «15,,09» is not a price; spaces are evened out.
    assert [(x.text, x.price, x.barcode) for x in lines] == [
        ("AFNAN 9 PM (U) edp 100ML", Decimal("13860"), "6290171002338"),
        ("X", Decimal("500"), None),
    ]
    with pytest.raises(PriceListError, match="не прайс 1С"):
        read_price_list(PriceFormat.usd_1c, erik_file([(None, "X", 500)]), "x.xlsx")
    with pytest.raises(PriceListError, match="не прайс со штрихкодами"):
        read_price_list(PriceFormat.barcode_kzt, kanat_file([("X", 500)]), "x.xlsx")


def test_suppliers_list(client, auth, suppliers):
    r = client.get("/api/admin/suppliers", headers=auth(UserRole.admin))
    assert r.status_code == 200
    assert [(s["name"], s["is_main"], s["price_format"]) for s in r.json()] == [
        ("Канат", True, "usd_1c"),
        ("Эрик", False, "barcode_kzt"),
    ]
    assert client.get("/api/admin/suppliers", headers=auth(UserRole.wholesale)).status_code == 403


def test_a_price_list_checks_then_applies_and_never_lowers_prices(
    client, auth, db, catalog, suppliers
):
    h = auth(UserRole.admin)
    kanat = suppliers["kanat"]
    coco50, coco100 = catalog["coco50"], catalog["coco100"]
    _priced(db, coco50, "10000", "16000", "13000", "12000")  # 60 / 30 / 20 %
    coco100.retail_price, coco100.wholesale_price = Decimal("30000"), None
    db.commit()  # no cost yet
    _link(client, h, kanat, "CHANEL Coco Mademoiselle L EdP 50", coco50.id)
    _link(client, h, kanat, "CHANEL Coco Mademoiselle L EdP 100", coco100.id)
    _link(client, h, kanat, "Флаконы КАПЛЯ 10 ml", None)

    price_list = kanat_file(
        [
            ("CHANEL  Coco Mademoiselle L EdP 50", 11000),  # spaces differ: the same line
            ("CHANEL Coco Mademoiselle L EdP 100", 20000),
            ("Флаконы КАПЛЯ 10 ml", 158.4),
            ("Tom Ford Oud Wood U EdP 50", 90000),
        ]
    )
    report = _upload(client, h, kanat, price_list)
    assert report["dry_run"] is True
    assert {k: report[k] for k in ("lines", "matched", "ignored", "by_barcode")} == {
        "lines": 4,
        "matched": 2,
        "ignored": 1,
        "by_barcode": 0,
    }
    assert [u["line"] for u in report["unmatched"]] == ["Tom Ford Oud Wood U EdP 50"]
    assert {k: report[k] for k in ("cost_up", "prices_raised", "remembered", "supplier_new")} == {
        "cost_up": 1,
        "prices_raised": 1,
        "remembered": 1,
        "supplier_new": 2,
    }
    first = report["changes"][0]
    assert (first["label"], first["retail_old"], first["retail_new"]) == (
        "Chanel Coco Mademoiselle, 50 мл",
        16000,
        17600,
    )
    assert _state(db, coco50.id) == (10000, 16000, 13000, 12000)  # a check changes nothing

    _upload(client, h, kanat, price_list, apply=True)
    assert _state(db, coco50.id) == (11000, 17600, 14300, 13200)
    # The first cost: the prices stay, their markups are remembered.
    v100 = db.get(ProductVariant, coco100.id)
    assert (v100.cost_price, v100.retail_price, v100.retail_markup) == (20000, 30000, 50)

    # A lower cost: shown, but the prices stay.
    report = _upload(
        client, h, kanat, kanat_file([("CHANEL Coco Mademoiselle L EdP 50", 9000)]), apply=True
    )
    assert (report["cost_down"], report["prices_kept"], report["prices_raised"]) == (1, 1, 0)
    assert _state(db, coco50.id) == (9000, 17600, 14300, 13200)
    # The markups are still those the prices were set with: a cost of 12 000 gives 19 200.
    _upload(client, h, kanat, kanat_file([("CHANEL Coco Mademoiselle L EdP 50", 12000)]), True)
    assert _state(db, coco50.id) == (12000, 19200, 15600, 14400)

    listed = {s["name"]: s for s in client.get("/api/admin/suppliers", headers=h).json()}
    assert (listed["Канат"]["prices"], listed["Канат"]["links"], listed["Канат"]["ignored"]) == (
        2,
        2,
        1,
    )
    assert listed["Канат"]["last_list"]["matched"] == 1


def test_the_main_supplier_sets_the_cost_the_other_fills_in(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    kanat, erik = suppliers["kanat"], suppliers["erik"]
    sauvage = catalog["sauvage100"]
    _priced(db, sauvage, "10000", "16000")
    db.add(VariantBarcode(variant_id=sauvage.id, code="0348901250146"))
    db.commit()

    # Before Канат's first list, Эрик's prices are kept but the cost waits: otherwise his price
    # would stand for every volume Канат also has.
    report = _upload(client, h, erik, erik_file([("348901250146", "DIOR SAUVAGE EDT 100", 11000)]))
    assert (report["matched"], report["by_barcode"], report["waiting_for_main"]) == (1, 1, True)
    assert (report["supplier_new"], report["cost_up"], report["prices_raised"]) == (1, 0, 0)

    # Канат's list is in and doesn't have it: Эрик's price is the cost. Found by barcode, also
    # read as UPC-A (12 digits).
    _upload(client, h, kanat, kanat_file([("Something else", 5000)]), apply=True)
    report = _upload(client, h, erik, erik_file([("0348901250146", "DIOR SAUVAGE EDT 100", 11000)]))
    assert report["waiting_for_main"] is False
    _upload(client, h, erik, erik_file([("0348901250146", "DIOR SAUVAGE EDT 100", 11000)]), True)
    assert _state(db, sauvage.id)[:2] == (11000, 17600)

    # Канат has it too: his price is the cost, even when Эрик is cheaper.
    _link(client, h, kanat, "DIOR Sauvage M EdT 100", sauvage.id)
    _upload(client, h, kanat, kanat_file([("DIOR Sauvage M EdT 100", 12000)]), apply=True)
    assert _state(db, sauvage.id)[:2] == (12000, 19200)
    report = _upload(
        client, h, erik, erik_file([("0348901250146", "DIOR SAUVAGE EDT 100", 9000)]), apply=True
    )
    assert (report["supplier_down"], report["cost_down"], report["prices_raised"]) == (1, 0, 0)
    assert _state(db, sauvage.id)[:2] == (12000, 19200)

    prices = client.get(f"/api/admin/products/{catalog['sauvage'].id}/supplier-prices", headers=h)
    assert [(p["supplier"], p["is_main"], p["price"], p["latest"]) for p in prices.json()] == [
        ("Канат", True, 12000, True),
        ("Эрик", False, 9000, True),
    ]
    # Not in Канат's next list: the price stays, marked as not from the latest list.
    _link(client, h, kanat, "Other", catalog["coco50"].id)
    _upload(client, h, kanat, kanat_file([("Other", 5000)]), apply=True)
    prices = client.get(f"/api/admin/products/{catalog['sauvage'].id}/supplier-prices", headers=h)
    assert prices.json()[0]["latest"] is False


def test_two_lines_for_one_volume_take_the_higher_price(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    kanat = suppliers["kanat"]
    _link(client, h, kanat, "A", catalog["coco50"].id)
    _link(client, h, kanat, "B", catalog["coco50"].id)
    report = _upload(client, h, kanat, kanat_file([("A", 9000), ("B", 9500)]), apply=True)
    assert (report["matched"], report["duplicates"]) == (2, 1)
    db.expire_all()
    stored = db.query(SupplierPrice).one()
    assert (stored.price, stored.line) == (9500, "B")


def test_wrong_file(client, auth, suppliers):
    h = auth(UserRole.admin)
    files = {"file": ("x.xlsx", erik_file([(None, "X", 500)]), "application/octet-stream")}
    r = client.post(
        f"/api/admin/suppliers/{suppliers['kanat'].id}/price-list", files=files, headers=h
    )
    assert r.status_code == 422
    assert "не прайс 1С" in r.json()["detail"]
    files = {"file": ("x.pdf", b"%PDF", "application/pdf")}
    r = client.post(
        f"/api/admin/suppliers/{suppliers['kanat'].id}/price-list", files=files, headers=h
    )
    assert r.status_code == 415


def test_links_from_a_file_and_back(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    kanat = suppliers["kanat"]
    rows = [
        ["Строка прайса", "Бренд", "Название", "Тип", "Объём, мл", "Тестер", "Не наш товар"],
        [
            "CHANEL Coco Mademoiselle L EdP 50",
            "chanel",
            "Coco  Mademoiselle",
            "EDP",
            50,
            "нет",
            None,
        ],
        ["DIOR Sauvage M EdT 100", "Dior", "Sauvage", "Туалетная вода", 100, None, None],
        ["Флаконы КАПЛЯ 10 ml", None, None, None, None, None, "да"],
        ["Unknown line", "Nobody", "Nothing", None, 50, None, None],
        ["Coco 30", "Chanel", "Coco Mademoiselle", "EDP", 30, "нет", None],
    ]
    files = {"file": ("links.xlsx", _xlsx(rows), "application/octet-stream")}
    r = client.post(f"/api/admin/suppliers/{kanat.id}/links/import", files=files, headers=h)
    assert r.status_code == 200, r.text
    report = r.json()
    assert (report["rows"], report["linked"], report["ignored"]) == (5, 2, 1)
    assert [(p["line"], p["error"]) for p in report["problems"]] == [
        ("Unknown line", "нет такого товара на сайте"),
        ("Coco 30", "у товара нет такого объёма"),
    ]

    report = _upload(
        client,
        h,
        kanat,
        kanat_file([("CHANEL Coco Mademoiselle L EdP 50", 9000), ("Флаконы КАПЛЯ 10 ml", 158)]),
    )
    assert (report["matched"], report["ignored"], report["unmatched"]) == (1, 1, [])

    r = client.get(f"/api/admin/suppliers/{kanat.id}/links/export", headers=h)
    sheet = load_workbook(io.BytesIO(r.content)).active
    got = [list(row) for row in sheet.iter_rows(values_only=True)]
    assert got[0][:3] == ["Строка прайса", "Бренд", "Название"]
    assert sorted(got[1:]) == sorted(
        [
            ["CHANEL Coco Mademoiselle L EdP 50", "Chanel", "Coco Mademoiselle", "EDP", 50, "нет",
             None],
            ["DIOR Sauvage M EdT 100", "Dior", "Sauvage", "EDT", 100, "нет", None],
            ["Флаконы КАПЛЯ 10 ml", None, None, None, None, None, "да"],
        ]
    )  # fmt: skip

    # A line can be unlinked: it is unmatched again.
    r = client.delete(
        f"/api/admin/suppliers/{kanat.id}/links",
        params={"line": "флаконы  каПЛЯ 10 ml"},
        headers=h,
    )
    assert r.status_code == 204
    report = _upload(client, h, kanat, kanat_file([("Флаконы КАПЛЯ 10 ml", 158)]))
    assert [u["line"] for u in report["unmatched"]] == ["Флаконы КАПЛЯ 10 ml"]


def test_a_big_links_file_takes_a_few_queries(db, catalog, suppliers):
    """On the live site every query is a round trip to the database, and a supplier's links file
    has thousands of rows: the links are loaded once, not looked up row by row."""
    kanat = suppliers["kanat"]
    rows = [["Строка прайса", "Бренд", "Название", "Тип", "Объём, мл", "Тестер", "Не наш товар"]]
    rows += [[f"Флакон {i}", None, None, None, None, None, "да"] for i in range(500)]
    # The same line twice is one link.
    rows += [["Coco 50", "Chanel", "Coco Mademoiselle", "EDP", 50, None, None]] * 2
    content = _xlsx(rows)
    statements = []

    def count(conn, cursor, statement, *args):
        statements.append(statement)

    event.listen(engine, "before_cursor_execute", count)
    try:
        for _ in range(2):  # new links, then the same file again
            statements.clear()
            result = import_links(db, kanat, content, "links.xlsx")
            db.commit()
            assert (result.rows, result.linked, result.ignored) == (502, 2, 500)
            assert len(statements) < 15
    finally:
        event.remove(engine, "before_cursor_execute", count)
    links = db.scalars(select(SupplierLink).where(SupplierLink.supplier_id == kanat.id)).all()
    assert len(links) == 501
    assert [link.variant_id for link in links if link.line == "Coco 50"] == [catalog["coco50"].id]


def test_the_export_has_each_suppliers_price(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    kanat, erik = suppliers["kanat"], suppliers["erik"]
    coco50, coco100 = catalog["coco50"], catalog["coco100"]
    db.add_all(
        [
            SupplierPrice(supplier_id=kanat.id, variant_id=coco50.id, price=9000, line="a"),
            SupplierPrice(supplier_id=erik.id, variant_id=coco50.id, price=8500.5, line="b"),
            SupplierPrice(supplier_id=erik.id, variant_id=coco100.id, price=12000, line="c"),
        ]
    )
    db.commit()

    r = client.get("/api/admin/import/export", headers=h)
    header, *cells = [list(row) for row in load_workbook(io.BytesIO(r.content)).active.values]
    start = header.index("Цена Канат")
    assert header[start : start + 3] == ["Цена Канат", "Цена Эрик", "Себестоимость"]
    got = {(row[1], row[header.index("Объём, мл")]): row[start : start + 2] for row in cells}
    assert got[("Coco Mademoiselle", 50)] == [9000, 8500.5]
    assert got[("Coco Mademoiselle", 100)] == [None, 12000]
    assert got[("Sauvage", 100)] == [None, None]

    # Imported back, the file changes nothing: the suppliers' prices come from their lists only.
    files = {"file": ("catalog.xlsx", r.content, "application/octet-stream")}
    report = client.post("/api/admin/import/catalog", files=files, headers=h).json()
    assert report["errors"] == [] and report["unmapped_columns"] == []
    counts = ("products_created", "products_updated", "variants_created", "variants_updated")
    assert [report[k] for k in counts] == [0, 0, 0, 0]


def test_a_receipt_line_takes_the_receipts_suppliers_price(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    kanat, erik = suppliers["kanat"], suppliers["erik"]
    coco50, coco100 = catalog["coco50"], catalog["coco100"]
    _priced(db, coco100, "11000", "18000")
    db.add_all(
        [
            SupplierPrice(supplier_id=kanat.id, variant_id=coco50.id, price=9000, line="a"),
            SupplierPrice(supplier_id=erik.id, variant_id=coco50.id, price=8500, line="b"),
            SupplierPrice(supplier_id=erik.id, variant_id=coco100.id, price=12000, line="c"),
        ]
    )
    db.commit()

    def lines(supplier: str | None) -> list[dict]:
        rid = client.post("/api/admin/receipts", json={"supplier": supplier}, headers=h).json()[
            "id"
        ]
        for v in (coco50, coco100):
            body = {"variant_id": v.id, "quantity": 1}
            r = client.post(f"/api/admin/receipts/{rid}/lines", json=body, headers=h)
            assert r.status_code == 200, r.text
        return r.json()["items"]

    # The receipt's supplier, written any way: its price list's price.
    erik_lines = lines(" эрик ")
    assert [i["cost_price"] for i in erik_lines] == [8500, 12000]
    assert erik_lines[0]["supplier_prices"] == [
        {"supplier": "Канат", "price": 9000},
        {"supplier": "Эрик", "price": 8500},
    ]
    assert erik_lines[1]["supplier_prices"][0] == {"supplier": "Канат", "price": None}
    # No supplier: the main supplier's price; where it has none, the volume's cost as before.
    assert [i["cost_price"] for i in lines(None)] == [9000, 11000]
    # A supplier without price lists: the volume's cost as before (coco50 has none).
    assert [i["cost_price"] for i in lines("Алматы опт")] == [None, 11000]

    # Autocomplete: the suppliers with price lists first.
    names = client.get("/api/admin/receipts/suppliers", headers=h).json()
    assert names[:2] == ["Канат", "Эрик"] and "Алматы опт" in names and "эрик" not in names


def test_the_till_quote_has_the_cost(client, auth, db, catalog):
    h = auth(UserRole.admin)
    _priced(db, catalog["coco50"], "60", "100")
    body = {"items": [{"variant_id": catalog["coco50"].id, "quantity": 2, "discount_percent": 0}]}
    line = client.post("/api/admin/store/quote", json=body, headers=h).json()["lines"][0]
    assert (line["cost_price"], line["line_total"]) == (60, 200)


def test_a_receipt_leaves_the_suppliers_cost_alone(client, auth, db, catalog, suppliers):
    h = auth(UserRole.admin)
    coco50, coco100 = catalog["coco50"], catalog["coco100"]
    _priced(db, coco50, "10000", "16000")
    coco100.retail_price, coco100.wholesale_price = Decimal("12000"), None
    db.add(SupplierPrice(supplier_id=suppliers["kanat"].id, variant_id=coco50.id, price=10000,
                         line="CHANEL Coco 50"))  # fmt: skip
    db.commit()
    rid = client.post("/api/admin/receipts", json={}, headers=h).json()["id"]
    for variant in (coco50, coco100):
        r = client.post(
            f"/api/admin/receipts/{rid}/lines",
            json={"variant_id": variant.id, "quantity": 10, "cost_price": 8000},
            headers=h,
        )
        assert r.status_code in (200, 201), r.text
    assert client.post(f"/api/admin/receipts/{rid}/post", headers=h).status_code == 200
    db.expire_all()
    # The price list sets the cost of a volume a supplier has.
    assert db.get(ProductVariant, coco50.id).cost_price == 10000
    # A volume no supplier has: the receipt gives it its first cost and the markups of its prices.
    v100 = db.get(ProductVariant, coco100.id)
    assert (v100.cost_price, v100.retail_markup) == (8000, 50)
