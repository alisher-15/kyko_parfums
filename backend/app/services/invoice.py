"""Delivery note (накладная) of an order as an Excel file, printed and put in the parcel.

The printable page with the same content (for PDF) is /admin/orders/<id>/invoice in the frontend.
"""

import io
from datetime import datetime
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.models import Gender, Order, OrderChannel, OrderItem, Product

SHOP_NAME = "Kyko Parfum"
PAYMENT_LABELS = {
    "cash": "Наличные",
    "card": "Карта",
    "transfer": "Kaspi / перевод",
    "other": "Другое",
}


GENDER_LABELS = {Gender.female: "Женский", Gender.male: "Мужской", Gender.unisex: "Унисекс"}


def product_details(
    db: Session, items: list[OrderItem]
) -> dict[int, tuple[str | None, Gender | None]]:
    """Type (EDP, EDT, ...) and gender of the products in the order, by product id.

    Order lines keep brand, name and volume, but not these: they are read from the catalog
    (a product deleted since then has none).
    """
    ids = {i.product_id for i in items if i.product_id is not None}
    if not ids:
        return {}
    rows = db.execute(select(Product.id, Product.type, Product.gender).where(Product.id.in_(ids)))
    return {pid: (ptype, gender) for pid, ptype, gender in rows}


def invoice_filename(order: Order) -> str:
    return f"nakladnaya_{order.id}.xlsx"


def _local(dt: datetime) -> datetime:
    return dt.astimezone(ZoneInfo(get_settings().timezone))


def build_invoice(order: Order, details: dict[int, tuple[str | None, Gender | None]]) -> bytes:
    """Lines removed before delivery (quantity 0) are left out; returns don't change the note."""
    settings = get_settings()
    money_fmt = f'#,##0 "{settings.currency_sign}"'
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Накладная {order.id}"

    ws.append([f"Накладная № {order.id} от {_local(order.created_at):%d.%m.%Y}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    info = [
        ("Продавец", SHOP_NAME),
        (
            "Покупатель",
            order.contact_name
            or ("Покупатель в магазине" if order.channel == OrderChannel.store else ""),
        ),
        ("Телефон", order.contact_phone),
        ("Email", order.contact_email),
        (
            "Адрес доставки",
            ", ".join(p for p in (order.delivery_city, order.delivery_address) if p),
        ),
        (
            "Оплата",
            PAYMENT_LABELS.get(order.payment_method.value) if order.payment_method else None,
        ),
        ("Комментарий", order.comment),
    ]
    for label, value in info:
        if value:
            ws.append([label, value])
            ws.cell(ws.max_row, 1).font = bold
    ws.append([])

    header = [
        "№", "Бренд", "Товар", "Тип", "Пол", "Объём, мл", "Тестер", "Артикул", "Кол-во", "Цена",
        "Сумма",
    ]  # fmt: skip
    qty_col, price_col, sum_col = 9, 10, 11
    ws.append(header)
    head_row = ws.max_row
    for col in range(1, len(header) + 1):
        c = ws.cell(head_row, col)
        c.font = bold
        c.border = box
        c.alignment = Alignment(horizontal="center", vertical="center")

    lines = [i for i in order.items if i.quantity > 0]
    for n, item in enumerate(lines, 1):
        ptype, gender = details.get(item.product_id, (None, None))
        ws.append(
            [
                n,
                item.brand_name,
                item.product_name,
                ptype,
                GENDER_LABELS.get(gender) if gender else None,
                item.volume_ml,
                "Да" if item.is_tester else "Нет",
                item.variant.sku if item.variant else None,
                item.quantity,
                float(item.price_applied),
                float(item.line_total),
            ]
        )
        row = ws.max_row
        for col in range(1, len(header) + 1):
            ws.cell(row, col).border = box
        for col in (4, 5, 6, 7, qty_col):
            ws.cell(row, col).alignment = Alignment(horizontal="center")
        ws.cell(row, price_col).number_format = money_fmt
        ws.cell(row, sum_col).number_format = money_fmt

    total_row = ws.max_row + 1
    ws.cell(total_row, qty_col - 1, "Итого").font = bold
    ws.cell(total_row, qty_col, sum(i.quantity for i in lines)).font = bold
    total = ws.cell(total_row, sum_col, float(sum(i.line_total for i in lines)))
    total.font = bold
    total.number_format = money_fmt

    ws.append([])
    ws.append([])
    ws.append(["Отпустил: ____________________", None, None, None, None, None, None,
               "Получил: ____________________"])  # fmt: skip

    for col, width in enumerate([5, 16, 30, 9, 11, 10, 8, 14, 8, 12, 13], 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
